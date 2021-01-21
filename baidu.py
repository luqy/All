from docx import Document
doc=Document('测试报告1.docx')
#print(doc)
table=doc.tables
for i in table:
    print(i.rows[14].cells[3].text)
    # print(len(i.rows))
    # print(len(i.columns))
    # for x in range(len(i.rows)):
    #     for y in range(len(y.columns)):
    #         print(i.cell(y,x).text)

